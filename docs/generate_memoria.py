# -*- coding: utf-8 -*-
"""Genera la memoria escrita del proyecto MeRenta en formato Word (.docx)."""

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

ACCENT = RGBColor(0x1F, 0x6F, 0xEB)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x55, 0x55, 0x55)

doc = Document()

# Márgenes compactos
for section in doc.sections:
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

# Estilo base
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.font.color.rgb = DARK
normal.paragraph_format.space_after = Pt(4)


def heading(text, level=1):
    h = doc.add_heading(level=level)
    h.paragraph_format.space_before = Pt(6)
    h.paragraph_format.space_after = Pt(3)
    run = h.add_run(text)
    run.font.color.rgb = ACCENT if level == 1 else DARK
    run.font.name = "Calibri"
    run.font.size = Pt(14)
    return h


def para(text="", bold=False, italic=False, size=10.5, color=DARK, align=None, space_after=4):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.color.rgb = color
    p.paragraph_format.space_after = Pt(space_after)
    return p


def bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
    p.add_run(text)
    for r in p.runs:
        r.font.size = Pt(10.5)
    return p


def link_line(label, url):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label)
    r.bold = True
    r.font.size = Pt(10.5)
    p.add_run(url).font.size = Pt(10.5)
    return p


def add_table(headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(9)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(val)
            run.font.size = Pt(9)
    return table


# ------------------------------------------------------------------
# 1. PORTADA
# ------------------------------------------------------------------
for _ in range(4):
    doc.add_paragraph()

para("MeRenta", bold=True, size=40, color=ACCENT, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para("Marketplace de alquiler entre particulares (P2P)", italic=True, size=15,
     color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=30)

para("Memoria del Proyecto Final", bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para("Ingeniería del Software II · Curso 2025–2026", size=12, color=GREY,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para("Universidad de León", size=12, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=36)

para("Equipo de desarrollo", bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
for name in [
    "Jose Ángel Mestas Díaz",
    "Elena Ondicol García",
    "Diego Pérez González",
    "Lucía González Rodríguez",
]:
    para(name, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)

doc.add_page_break()

# ------------------------------------------------------------------
# 2. INTRODUCCIÓN
# ------------------------------------------------------------------
heading("1. Introducción", 1)
para("MeRenta es una aplicación web full-stack que implementa un marketplace de alquiler "
     "entre particulares, en la línea de Wallapop pero orientado al alquiler temporal de "
     "productos en lugar de a su compraventa. Cualquier usuario puede publicar sus objetos "
     "para que otros los alquilen y gestionar todo el ciclo desde la plataforma.")
para("Problema y público objetivo. ", bold=True, space_after=0)
para("Muchos objetos de uso esporádico permanecen infrautilizados. MeRenta conecta a quien "
     "los posee con quien los necesita puntualmente, aportando confianza mediante pagos "
     "seguros, un seguro obligatorio por reserva, valoraciones y mensajería interna. Va "
     "dirigido a particulares que quieren rentabilizar objetos o alquilarlos temporalmente.")
para("Objetivo. ", bold=True, space_after=0)
para("Desarrollar y desplegar en producción un marketplace completo que cubra el flujo de "
     "extremo a extremo: autenticación, publicación y búsqueda, reservas con máquina de "
     "estados, pagos con Stripe, chat en tiempo real, valoraciones y panel de administración.")

# ------------------------------------------------------------------
# 3. EQUIPO DE TRABAJO
# ------------------------------------------------------------------
heading("2. Equipo de Trabajo", 1)
para("El equipo lo forman cuatro integrantes que participaron de forma equitativa en todas "
     "las áreas (backend, frontend, base de datos, pruebas y despliegue), como refleja el "
     "historial de commits y Pull Requests. No hubo división rígida de roles: todos "
     "trabajaron tanto en Go como en React según las prioridades de cada iteración.")
bullet("cada funcionalidad o corrección fue una tarea en GitHub Projects, con su rama de "
       "corta duración e issue.", bold_prefix="Organización por tareas: ")
bullet("cada Pull Request fue revisado por otro miembro antes de fusionar.",
       bold_prefix="Revisión entre compañeros: ")
bullet("las decisiones clave se documentaron como ADRs en /docs/adr.",
       bold_prefix="Decisiones documentadas: ")

# ------------------------------------------------------------------
# 4. REQUISITOS Y ORGANIZACIÓN
# ------------------------------------------------------------------
heading("3. Requisitos y Organización del Proyecto", 1)
para("Requisitos funcionales principales:", bold=True, space_after=2)
for t in [
    "Registro e inicio de sesión con JWT y roles (user / admin).",
    "Publicación, edición y eliminación de productos con imágenes y condiciones de uso.",
    "Búsqueda y filtrado de productos y favoritos.",
    "Reservas gestionadas mediante una máquina de estados.",
    "Pago con Stripe y seguro obligatorio por reserva.",
    "Chat en tiempo real (WebSockets) y sistema de valoraciones.",
    "Panel de administración: usuarios, pagos, incidencias, verificación y ajustes.",
]:
    bullet(t)
para("Gestión de tareas. ", bold=True, space_after=0)
para("Se utilizó GitHub Projects con un tablero Kanban; cada issue se asignó a un "
     "responsable, lo que permite verificar la contribución individual.")
link_line("Tablero: ", "https://github.com/orgs/isw2-unileon/projects/13")

# ------------------------------------------------------------------
# 5. DISEÑO Y ARQUITECTURA
# ------------------------------------------------------------------
heading("4. Diseño y Arquitectura", 1)
para("MeRenta sigue una arquitectura cliente–servidor de tres capas, organizada como "
     "monorepo. El frontend es una SPA que consume una API REST; el backend concentra la "
     "lógica de negocio; y la persistencia recae en PostgreSQL (Supabase). Toda respuesta de "
     "la API usa el envoltorio { success, data, error }.")
para("Navegador (React + TS)  ⇄  API REST (Go + Gin)  ⇄  PostgreSQL (Supabase)",
     align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, color=GREY, space_after=0)
para("Servicios integrados: Stripe (pagos) y WebSockets (chat).",
     align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, size=9, color=GREY, space_after=4)
add_table(
    ["Capa", "Tecnología", "Justificación"],
    [
        ["Backend", "Go (Gin)", "Obligatorio; rendimiento y concurrencia para WebSockets."],
        ["Frontend", "React + TS + Vite + Tailwind", "Tipado estático y desarrollo ágil de UI."],
        ["Base de datos", "PostgreSQL + PostGIS (Supabase)", "Datos relacionales y consultas geográficas."],
        ["Acceso a datos", "sqlc", "Código Go tipado generado desde SQL, sin ORM."],
        ["Autenticación", "JWT (cookie HttpOnly)", "Sesiones seguras y stateless."],
        ["Pagos / Tiempo real", "Stripe / WebSockets", "Pasarela estándar y chat bidireccional."],
        ["Infraestructura", "Docker, Render, GitHub Actions", "Contenedores, despliegue y CI/CD."],
    ],
)
para("")
para("El backend se estructura en capas (handler → service → sqlcdb) con inyección de "
     "dependencias manual; el frontend separa lógica y vista (ficheros *.logic.ts) para "
     "facilitar las pruebas.")

# ------------------------------------------------------------------
# 6. DESARROLLO E IMPLEMENTACIÓN
# ------------------------------------------------------------------
heading("5. Desarrollo e Implementación", 1)
para("Buenas prácticas:", bold=True, space_after=2)
bullet("ramas de corta duración por tarea integradas con frecuencia vía Pull Request "
       "(Trunk Based Development); dev sirve solo de puente del código estable a main.",
       bold_prefix="Flujo: ")
bullet("Pull Requests revisados con comentarios trazables; arquitectura por capas y "
       "separación de responsabilidades.", bold_prefix="Revisión y diseño: ")
bullet("golangci-lint (gosec, errcheck, funlen) y ESLint + Prettier, verificados en CI; "
       "Conventional Commits e inglés en todo el repositorio.",
       bold_prefix="Calidad y convenciones: ")
para("Ejemplo (reserva con pago y seguro). ", bold=True, space_after=0)
para("Atraviesa todas las capas e integra un servicio externo: el usuario elige fechas en "
     "BookingCard (con validación de disponibilidad y días mín./máx.), pasa al checkout donde "
     "se procesa el pago con Stripe y se asocia el seguro obligatorio, y al confirmarse el "
     "pago la reserva avanza en su máquina de estados. El cálculo se aisló en *.logic.ts con "
     "tests y el backend valida y persiste vía la capa de servicios.")
para("Pruebas (tres niveles). ", bold=True, space_after=0)
para("Unitarias: Go con -race (45 ficheros) y Vitest (26 ficheros), con umbral de cobertura "
     "del 50 % en CI. De integración: script dedicado en el pipeline de backend. E2E: "
     "Playwright sobre los flujos críticos (auth, reservas, checkout, marketplace, perfil).")

# ------------------------------------------------------------------
# 7. CI/CD Y DESPLIEGUE
# ------------------------------------------------------------------
heading("6. CI/CD y Despliegue", 1)
para("La integración continua usa GitHub Actions con tres workflows filtrados por ruta, de "
     "modo que solo se ejecuta lo afectado por cada cambio:")
bullet("lint → format-check → tests unitarios → tests de integración → build.",
       bold_prefix="backend.yml: ")
bullet("lint → format-check → tests → build.", bold_prefix="frontend.yml: ")
bullet("pruebas Playwright (disparo manual).", bold_prefix="e2e.yml: ")
para("El build depende de que pasen las fases previas y se publican artefactos de cobertura y "
     "compilación. La aplicación está desplegada en Render.com con la base de datos en "
     "Supabase, con entornos de producción y desarrollo.")
link_line("Producción: ", "https://merenta.onrender.com/")
link_line("Desarrollo: ", "https://merenta-p9ny.onrender.com/home")

# ------------------------------------------------------------------
# 8. RESULTADOS FINALES
# ------------------------------------------------------------------
heading("7. Resultados Finales", 1)
para("El producto está operativo en producción. Funcionalidades implementadas: landing y "
     "páginas legales; registro/login con sesión persistente; gestión de productos con "
     "condiciones de uso; búsqueda, filtrado y favoritos; reservas con máquina de estados, "
     "checkout con Stripe y seguro obligatorio; chat en tiempo real (con cifrado y búsqueda); "
     "perfiles con valoraciones; y panel de administración (usuarios, pagos, incidencias y "
     "verificación).")
link_line("Producto desplegado: ", "https://merenta.onrender.com/")

# ------------------------------------------------------------------
# 9. CONCLUSIONES
# ------------------------------------------------------------------
heading("8. Conclusiones y Lecciones Aprendidas", 1)
para("Aprendimos a construir y desplegar un producto real de principio a fin: pagos con "
     "Stripe, tiempo real con WebSockets, organización en monorepo, SQL tipado con sqlc y un "
     "pipeline de CI/CD con pruebas en tres niveles. La mayor dificultad fue la autenticación "
     "de las conexiones WebSocket en producción, que requirió varias iteraciones para "
     "estabilizar el chat, además de afinar pagos y CORS en el entorno desplegado; todo se "
     "resolvió mediante ramas de corrección y revisión en Pull Requests. Con más tiempo "
     "ampliaríamos la cobertura de tests por encima del 50 %, añadiríamos más pruebas E2E y "
     "automatizaríamos por completo el despliegue del entorno E2E.")

# ------------------------------------------------------------------
# 10. ENLACES
# ------------------------------------------------------------------
heading("9. Enlaces Relevantes", 1)
for label, url in [
    ("Repositorio GitHub: ", "https://github.com/isw2-unileon/MeRenta"),
    ("Despliegue (producción): ", "https://merenta.onrender.com/"),
    ("Despliegue (desarrollo): ", "https://merenta-p9ny.onrender.com/home"),
    ("GitHub Projects: ", "https://github.com/orgs/isw2-unileon/projects/13"),
    ("Panel de Render: ", "https://dashboard.render.com/project/prj-d6spb24hg0os73fcdhn0"),
    ("Base de datos (Supabase): ", "https://supabase.com/dashboard/project/dbmingfeohxjfzxelasn"),
]:
    link_line(label, url)

out = "docs/Memoria_MeRenta_v2.docx"
doc.save(out)
print("Guardado:", out)
